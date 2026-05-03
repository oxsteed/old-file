from __future__ import annotations

import base64
import html
import io
import json
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import qrcode
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT / "docs" / "marketing" / "flyer"
HTML_OUTPUT = OUTPUT_DIR / "oxsteed_flyer_v3_larger_font_true.html"
PDF_OUTPUT = OUTPUT_DIR / "oxsteed_flyer.pdf"
DOCX_OUTPUT = OUTPUT_DIR / "oxsteed_flyer.docx"
QR_OUTPUT = OUTPUT_DIR / "oxsteed_flyer_qr.png"
LOGO_PNG_OUTPUT = OUTPUT_DIR / "oxsteed_logo.png"
LOGO_SVG_PATH = REPO_ROOT / "client" / "public" / "icons" / "icon.svg"
LOCALE_PATH = REPO_ROOT / "client" / "src" / "locales" / "en.json"

CTA_URL = "https://oxsteed.com/consulting-intake"
CTA_LABEL = "oxsteed.com/consulting-intake"

FEATURE_CARDS = [
    {
        "title": "Launch ready",
        "body": (
            "Move from idea to a client-ready business with setup, systems, and next steps."
        ),
    },
    {
        "title": "AI-powered",
        "body": (
            "Use AI-assisted research, positioning, and planning to make decisions faster."
        ),
    },
    {
        "title": "Full-stack tech",
        "body": (
            "Get practical support for websites, forms, booking, payments, and automation."
        ),
    },
]

HOW_IT_WORKS_STEPS = [
    {
        "step": "01",
        "title": "Apply",
        "body": "Submit the intake form and tell us what you are building.",
    },
    {
        "step": "02",
        "title": "Kickoff call",
        "body": "We review your goals, blockers, and the right package.",
    },
    {
        "step": "03",
        "title": "Build the foundation",
        "body": "We set up the brand, identity, workflows, website, and tools you need.",
    },
    {
        "step": "04",
        "title": "Launch with clarity",
        "body": "You leave with assets, guidance, and a clear launch path.",
    },
]

WHY_OXSTEED_POINTS = [
    "Real implementation help, not theory alone.",
    "Built for trades, services, and local operators.",
    "Clear packages and practical deliverables.",
    "One partner for strategy, systems, and launch support.",
]


def _load_locale() -> dict[str, Any]:
    return json.loads(LOCALE_PATH.read_text(encoding="utf-8"))


def _load_package_details() -> list[dict[str, Any]]:
    package_details = (
        _load_locale()["consulting"]["intake"]["packageDetails"]
    )
    ordered_keys = ["initial", "foundation", "launch"]
    packages: list[dict[str, Any]] = []

    for key in ordered_keys:
        details = package_details[key]
        raw_features = details["features"]
        if key == "initial":
            highlights = raw_features[:3]
        elif key == "foundation":
            highlights = raw_features[:3]
        else:
            highlights = raw_features[:4]

        packages.append(
            {
                "key": key,
                "name": details["name"],
                "price": details["price"],
                "timeline": details["timeline"],
                "for_who": details["forWho"],
                "highlights": highlights,
                "popular": details.get("popular", False),
            }
        )

    return packages


PACKAGES = _load_package_details()


def _make_qr_png_bytes(target_url: str) -> bytes:
    qr = qrcode.QRCode(
        version=None,
        box_size=8,
        border=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
    )
    qr.add_data(target_url)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _qr_data_uri(target_url: str) -> str:
    encoded = base64.b64encode(_make_qr_png_bytes(target_url)).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def _logo_svg_markup() -> str:
    return LOGO_SVG_PATH.read_text(encoding="utf-8")


def build_context() -> dict[str, Any]:
    return {
        "brand": "OxSteed Consulting",
        "tagline": "Launch support for builders, operators, trades, and service businesses.",
        "subtitle": (
            "Business launch support for people building a serious service business."
        ),
        "eyebrow": "Business launch support",
        "feature_cards": FEATURE_CARDS,
        "how_it_works": HOW_IT_WORKS_STEPS,
        "packages": PACKAGES,
        "why_oxsteed": WHY_OXSTEED_POINTS,
        "cta_title": "Ready to launch?",
        "cta_body": (
            "Scan the QR code or visit the link to start your intake form and "
            "book the right next step."
        ),
        "cta_url": CTA_URL,
        "cta_label": CTA_LABEL,
        "logo_svg": _logo_svg_markup(),
        "qr_data_uri": _qr_data_uri(CTA_URL),
    }


def _render_feature_cards(cards: list[dict[str, str]]) -> str:
    rendered: list[str] = []
    for card in cards:
        rendered.append(
            f"""
            <article class="mini-card">
              <div class="mini-card__cap"></div>
              <h3>{html.escape(card["title"])}</h3>
              <p>{html.escape(card["body"])}</p>
            </article>
            """
        )
    return "".join(rendered)


def _render_steps(steps: list[dict[str, str]]) -> str:
    rendered: list[str] = []
    for step in steps:
        rendered.append(
            f"""
            <article class="step-card">
              <div class="step-card__number">{html.escape(step["step"])}</div>
              <h3>{html.escape(step["title"])}</h3>
              <p>{html.escape(step["body"])}</p>
            </article>
            """
        )
    return "".join(rendered)


def _render_packages(packages: list[dict[str, Any]]) -> str:
    rendered: list[str] = []
    for package in packages:
        badge = (
            '<div class="package-card__badge">Most popular</div>'
            if package["popular"]
            else ""
        )
        highlights = "".join(
            f"<li>{html.escape(feature)}</li>" for feature in package["highlights"]
        )
        rendered.append(
            f"""
            <article class="package-card{' package-card--popular' if package['popular'] else ''}">
              {badge}
              <div class="package-card__price">{html.escape(package["price"])}</div>
              <h3>{html.escape(package["name"])}</h3>
              <p class="package-card__timeline">{html.escape(package["timeline"])}</p>
              <p class="package-card__for-who">{html.escape(package["for_who"])}</p>
              <ul>{highlights}</ul>
            </article>
            """
        )
    return "".join(rendered)


def _render_why_points(points: list[str]) -> str:
    return "".join(f"<li>{html.escape(point)}</li>" for point in points)


def _render_offer_rows(cards: list[dict[str, str]]) -> str:
    rendered: list[str] = []
    for card in cards:
        rendered.append(
            f"""
            <div class="print-list-row">
              <strong>{html.escape(card["title"])}</strong>
              <span>{html.escape(card["body"])}</span>
            </div>
            """
        )
    return "".join(rendered)


def _render_step_rows(steps: list[dict[str, str]]) -> str:
    rendered: list[str] = []
    for step in steps:
        rendered.append(
            f"""
            <div class="print-list-row">
              <strong>{html.escape(step["step"])} {html.escape(step["title"])}</strong>
              <span>{html.escape(step["body"])}</span>
            </div>
            """
        )
    return "".join(rendered)


def _render_package_rows(packages: list[dict[str, Any]]) -> str:
    rendered: list[str] = []
    for package in packages:
        features = "; ".join(package["highlights"])
        rendered.append(
            f"""
            <div class="print-package-row">
              <div class="print-package-row__top">
                <strong>{html.escape(package["name"])}</strong>
                <span>{html.escape(package["price"])}</span>
              </div>
              <div class="print-package-row__meta">{html.escape(package["timeline"])} - {html.escape(package["for_who"])}</div>
              <div class="print-package-row__features">{html.escape(features)}</div>
            </div>
            """
        )
    return "".join(rendered)


def render_html(context: dict[str, Any]) -> str:
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>OxSteed Consulting Flyer</title>
  <style>
    :root {{
      --navy: #16213e;
      --navy-soft: #24304f;
      --orange: #f97316;
      --orange-deep: #ea580c;
      --cream: #fff8ef;
      --paper: #fffdf9;
      --ink: #172033;
      --muted: #55617b;
      --line: #d7deea;
      --shadow: 0 20px 50px rgba(22, 33, 62, 0.12);
      --radius: 22px;
    }}

    * {{
      box-sizing: border-box;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }}

    html, body {{
      margin: 0;
      padding: 0;
      background: #eef2f7;
      color: var(--ink);
      font-family: Arial, Helvetica, sans-serif;
      overflow: visible;
    }}

    body {{
      padding: 28px;
    }}

    .flyer {{
      max-width: 920px;
      margin: 0 auto;
    }}

    .page {{
      background: var(--paper);
      border-radius: 28px;
      box-shadow: var(--shadow);
      overflow: hidden;
    }}

    .page + .page {{
      margin-top: 18px;
    }}

    .hero {{
      padding: 24px 24px 18px;
      background:
        radial-gradient(circle at top right, rgba(249, 115, 22, 0.18), transparent 32%),
        linear-gradient(135deg, var(--navy) 0%, #0f172d 100%);
      color: white;
    }}

    .hero__header {{
      display: grid;
      grid-template-columns: minmax(0, 1fr) auto;
      gap: 20px;
      align-items: end;
    }}

    .brand {{
      display: flex;
      align-items: center;
      gap: 16px;
    }}

    .brand__logo {{
      width: 58px;
      height: 58px;
      flex: 0 0 58px;
      border-radius: 14px;
      background: rgba(255, 255, 255, 0.08);
      padding: 4px;
    }}

    .brand__text small {{
      display: inline-block;
      text-transform: uppercase;
      letter-spacing: 0.12em;
      font-size: 11px;
      font-weight: 700;
      color: #ffd5b5;
      margin-bottom: 8px;
    }}

    .brand__text h1 {{
      margin: 0 0 6px;
      font-size: 30px;
      line-height: 1.05;
    }}

    .brand__text p {{
      margin: 0;
      max-width: 54ch;
      font-size: 13px;
      line-height: 1.5;
      color: rgba(255, 255, 255, 0.88);
    }}

    .hero__meta {{
      text-align: right;
      align-self: start;
    }}

    .hero__pill {{
      display: inline-block;
      padding: 8px 12px;
      border-radius: 999px;
      background: rgba(255, 255, 255, 0.12);
      border: 1px solid rgba(255, 255, 255, 0.16);
      font-size: 12px;
      font-weight: 700;
      letter-spacing: 0.06em;
      text-transform: uppercase;
    }}

    .hero__subtitle {{
      margin-top: 14px;
      padding: 12px 14px;
      background: rgba(255, 255, 255, 0.08);
      border-radius: 18px;
      font-size: 12.5px;
      line-height: 1.45;
      color: rgba(255, 255, 255, 0.94);
    }}

    .content {{
      padding: 18px 22px 22px;
    }}

    .content--split {{
      display: grid;
      grid-template-columns: 1.1fr 0.9fr;
      gap: 16px;
      align-items: start;
    }}

    .stack {{
      display: grid;
      gap: 16px;
    }}

    .section {{
      margin-bottom: 18px;
    }}

    .section:last-child {{
      margin-bottom: 0;
    }}

    .section__heading {{
      margin: 0 0 10px;
      font-size: 17px;
      line-height: 1.15;
      color: var(--navy);
    }}

    .section__heading span {{
      color: var(--orange-deep);
    }}

    .section__note {{
      margin: 0 0 10px;
      color: var(--muted);
      font-size: 11.5px;
      line-height: 1.4;
    }}

    .offer-grid {{
      display: grid;
      gap: 10px;
    }}

    .grid-4 {{
      display: grid;
      gap: 10px;
    }}

    .mini-card,
    .step-card,
    .package-card {{
      border: 1px solid var(--line);
      border-radius: var(--radius);
      background: white;
      page-break-inside: avoid;
      break-inside: avoid;
    }}

    .mini-card {{
      padding: 12px 13px;
      background: linear-gradient(180deg, #ffffff 0%, #fff8f0 100%);
    }}

    .mini-card__cap {{
      width: 30px;
      height: 4px;
      border-radius: 999px;
      background: var(--orange);
      margin-bottom: 8px;
    }}

    .mini-card h3,
    .step-card h3,
    .package-card h3 {{
      margin: 0 0 6px;
      font-size: 14px;
      line-height: 1.2;
      color: var(--navy);
    }}

    .mini-card p,
    .step-card p,
    .package-card p,
    .package-card li {{
      margin: 0;
      color: var(--muted);
      font-size: 11px;
      line-height: 1.35;
    }}

    .step-card {{
      padding: 10px 12px;
      background: #fcfdff;
      display: grid;
      grid-template-columns: auto 1fr;
      gap: 10px;
      align-items: start;
    }}

    .step-card__number {{
      width: 28px;
      height: 28px;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      border-radius: 999px;
      background: var(--navy);
      color: white;
      font-size: 10px;
      font-weight: 700;
      margin-bottom: 0;
    }}

    .step-card__body {{
      min-width: 0;
    }}

    .package-grid {{
      display: grid;
      grid-template-columns: repeat(3, minmax(0, 1fr));
      gap: 10px;
      align-items: start;
    }}

    .package-card {{
      position: relative;
      padding: 12px;
      min-height: 100%;
    }}

    .package-card--popular {{
      border-color: rgba(249, 115, 22, 0.45);
      box-shadow: 0 10px 22px rgba(249, 115, 22, 0.12);
      background: linear-gradient(180deg, #fff9f3 0%, #ffffff 100%);
    }}

    .package-card__badge {{
      position: absolute;
      top: 12px;
      right: 12px;
      padding: 5px 8px;
      border-radius: 999px;
      background: rgba(249, 115, 22, 0.12);
      color: var(--orange-deep);
      font-size: 10px;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.08em;
    }}

    .package-card__price {{
      margin-bottom: 6px;
      font-size: 23px;
      font-weight: 800;
      color: var(--orange-deep);
      line-height: 1;
    }}

    .package-card__timeline {{
      margin-bottom: 6px !important;
      color: var(--navy-soft) !important;
      font-weight: 700;
    }}

    .package-card__for-who {{
      margin-bottom: 8px !important;
    }}

    .package-card ul {{
      margin: 0;
      padding-left: 16px;
    }}

    .package-card li + li {{
      margin-top: 4px;
    }}

    .why-list {{
      margin: 0;
      padding-left: 20px;
      columns: 1;
      column-gap: 0;
    }}

    .why-list li {{
      margin-bottom: 6px;
      color: var(--muted);
      font-size: 11.5px;
      line-height: 1.35;
      break-inside: avoid;
    }}

    .cta {{
      display: grid;
      grid-template-columns: minmax(0, 1fr) auto;
      gap: 14px;
      align-items: center;
      padding: 14px;
      border-radius: 20px;
      background: linear-gradient(135deg, #fff6ec 0%, #fff 100%);
      border: 1px solid rgba(249, 115, 22, 0.24);
    }}

    .cta h2 {{
      margin: 0 0 6px;
      font-size: 21px;
      line-height: 1.1;
      color: var(--navy);
    }}

    .cta p {{
      margin: 0 0 10px;
      font-size: 11.5px;
      line-height: 1.4;
      color: var(--muted);
      max-width: 50ch;
    }}

    .cta__link {{
      display: inline-flex;
      align-items: center;
      gap: 10px;
      padding: 10px 14px;
      border-radius: 999px;
      background: var(--orange);
      color: white;
      font-size: 12px;
      font-weight: 700;
      text-decoration: none;
      box-shadow: 0 10px 24px rgba(249, 115, 22, 0.22);
    }}

    .cta__url {{
      margin-top: 8px;
      font-size: 11px;
      color: var(--navy-soft);
    }}

    .cta__qr {{
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 8px;
      width: 132px;
      padding: 8px;
      border-radius: 18px;
      background: white;
      border: 1px solid var(--line);
    }}

    .cta__qr img {{
      width: 100px;
      height: 100px;
      display: block;
      border-radius: 10px;
    }}

    .cta__qr span {{
      font-size: 11px;
      color: var(--muted);
      text-transform: uppercase;
      letter-spacing: 0.08em;
      font-weight: 700;
    }}

    .services-section {{
      display: none;
    }}

    .screen-only {{
      display: block;
    }}

    .print-only {{
      display: none;
    }}

    @media (max-width: 820px) {{
      body {{
        padding: 12px;
      }}

      .hero__header,
      .cta {{
        grid-template-columns: 1fr;
      }}

      .hero__meta {{
        text-align: left;
      }}

      .content--split,
      .grid-4,
      .package-grid {{
        grid-template-columns: 1fr;
      }}

      .step-card {{
        grid-template-columns: auto 1fr;
      }}
    }}

    @page {{
      size: A4;
      margin: 15mm;
    }}

    @media print {{
      html, body {{
        background: white;
        padding: 0;
        margin: 0;
      }}

      .screen-only {{
        display: none !important;
      }}

      .print-only {{
        display: block !important;
      }}

      .page {{
        max-width: none;
        box-shadow: none;
        border-radius: 0;
        break-after: page;
        page-break-after: always;
      }}

      .page:last-child {{
        break-after: auto;
        page-break-after: auto;
      }}

      .services-section {{
        display: none !important;
      }}

      a {{
        text-decoration: none;
      }}

      .flyer {{
        max-width: none;
      }}

      .page {{
        overflow: visible;
      }}

      .hero {{
        padding: 16px 18px 12px;
      }}

      .brand__logo {{
        width: 48px;
        height: 48px;
        flex-basis: 48px;
      }}

      .brand__text h1 {{
        font-size: 24px;
      }}

      .brand__text p,
      .hero__subtitle {{
        font-size: 10.5px;
        line-height: 1.3;
      }}

      .hero__pill {{
        font-size: 10px;
        padding: 6px 10px;
      }}

      .content,
      .content--split {{
        display: block;
        padding: 14px 18px 16px;
      }}

      .section {{
        margin-bottom: 12px;
      }}

      .section__heading {{
        font-size: 14px;
        margin-bottom: 6px;
      }}

      .section__note {{
        font-size: 9.5px;
        margin-bottom: 6px;
      }}

      .print-list,
      .print-package-list {{
        display: grid;
        gap: 6px;
      }}

      .print-list-row,
      .print-package-row {{
        border: 1px solid var(--line);
        border-radius: 12px;
        padding: 6px 8px;
        background: white;
      }}

      .print-list-row strong,
      .print-package-row strong {{
        display: block;
        font-size: 10px;
        line-height: 1.2;
        color: var(--navy);
        margin-bottom: 2px;
      }}

      .print-list-row span,
      .print-package-row__meta,
      .print-package-row__features {{
        display: block;
        font-size: 8.8px;
        line-height: 1.25;
        color: var(--muted);
      }}

      .print-package-row__top {{
        display: flex;
        justify-content: space-between;
        gap: 8px;
        align-items: baseline;
      }}

      .print-package-row__top span {{
        font-size: 11px;
        font-weight: 700;
        color: var(--orange-deep);
        white-space: nowrap;
      }}

      .why-list {{
        padding-left: 16px;
      }}

      .why-list li {{
        font-size: 9px;
        line-height: 1.2;
        margin-bottom: 4px;
      }}

      .cta {{
        padding: 10px 12px;
        gap: 10px;
        grid-template-columns: minmax(0, 1fr) 92px;
      }}

      .cta h2 {{
        font-size: 17px;
        margin-bottom: 4px;
      }}

      .cta p,
      .cta__url {{
        font-size: 9px;
        line-height: 1.25;
      }}

      .cta__link {{
        padding: 8px 12px;
        font-size: 10px;
      }}

      .cta__qr {{
        width: 92px;
        padding: 6px;
        gap: 4px;
      }}

      .cta__qr img {{
        width: 72px;
        height: 72px;
      }}

      .cta__qr span {{
        font-size: 8px;
      }}
    }}
  </style>
</head>
<body>
  <main class="flyer">
    <section class="page">
      <section class="hero">
        <div class="hero__header">
          <div class="brand">
            <div class="brand__logo">{context["logo_svg"]}</div>
            <div class="brand__text">
              <small>{html.escape(context["eyebrow"])}</small>
              <h1>{html.escape(context["brand"])}</h1>
              <p>{html.escape(context["tagline"])}</p>
            </div>
          </div>
          <div class="hero__meta">
            <div class="hero__pill">Print-ready handout</div>
          </div>
        </div>
      </section>
      <div class="hero__subtitle">{html.escape(context["subtitle"])}</div>
      <section class="content content--split">
        <div class="stack">
          <section class="section screen-only">
            <h2 class="section__heading">What We Offer</h2>
            <div class="offer-grid">{_render_feature_cards(context["feature_cards"])}</div>
          </section>

          <section class="section print-only">
            <h2 class="section__heading">What We Offer</h2>
            <div class="print-list">{_render_offer_rows(context["feature_cards"])}</div>
          </section>

          <section class="section">
            <h2 class="section__heading">Why OxSteed?</h2>
            <ul class="why-list">{_render_why_points(context["why_oxsteed"])}</ul>
          </section>
        </div>

        <section class="section screen-only">
          <h2 class="section__heading">How It Works</h2>
          <div class="grid-4">{_render_steps(context["how_it_works"])}</div>
        </section>

        <section class="section print-only">
          <h2 class="section__heading">How It Works</h2>
          <div class="print-list">{_render_step_rows(context["how_it_works"])}</div>
        </section>
      </section>
    </section>

    <section class="page">
      <section class="content">
        <section class="section screen-only">
          <h2 class="section__heading">Consulting Packages</h2>
          <p class="section__note">
            Tight, shareable overview for the flyer version. The full service catalog stays off the print handout.
          </p>
          <div class="package-grid">{_render_packages(context["packages"])}</div>
        </section>

        <section class="section print-only">
          <h2 class="section__heading">Consulting Packages</h2>
          <p class="section__note">
            Tight, shareable overview for the flyer version. The full service catalog stays off the print handout.
          </p>
          <div class="print-package-list">{_render_package_rows(context["packages"])}</div>
        </section>

        <section class="section">
          <div class="cta">
            <div>
              <h2>{html.escape(context["cta_title"])}</h2>
              <p>{html.escape(context["cta_body"])}</p>
              <a class="cta__link" href="{html.escape(context["cta_url"])}">
                Apply now
              </a>
              <div class="cta__url">{html.escape(context["cta_label"])}</div>
            </div>
            <div class="cta__qr">
              <img src="{context["qr_data_uri"]}" alt="QR code for OxSteed consulting intake" />
              <span>Scan to apply</span>
            </div>
          </div>
        </section>
      </section>
    </section>
  </main>
</body>
</html>
"""


def _write_file(path: Path, contents: str | bytes) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(contents, bytes):
        path.write_bytes(contents)
    else:
        path.write_text(contents, encoding="utf-8")
    return path


def _chrome_path() -> str:
    direct_candidates = [
        "/opt/google/chrome/chrome",
        "/usr/bin/chromium",
        "/usr/bin/chromium-browser",
    ]
    for candidate in direct_candidates:
        if Path(candidate).exists():
            return candidate
    for candidate in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser"):
        resolved = shutil.which(candidate)
        if resolved:
            return resolved
    raise RuntimeError("Chrome/Chromium is required to export the flyer PDF.")


def _run_chrome(arguments: list[str]) -> None:
    with tempfile.TemporaryDirectory(prefix="oxsteed-flyer-chrome-") as temp_dir:
        command = [
            _chrome_path(),
            f"--user-data-dir={temp_dir}",
        ] + arguments
        subprocess.run(command, check=True, capture_output=True, text=True)


def _rasterize_logo_png(output_path: Path) -> Path:
    temp_html = output_path.with_name("_logo_raster_source.html")
    temp_html.write_text(
        (
            "<!doctype html><html><body style='margin:0;background:white;'>"
            f"<img src='{LOGO_SVG_PATH.resolve().as_uri()}' "
            "style='width:512px;height:512px;display:block' alt='OxSteed logo' />"
            "</body></html>"
        ),
        encoding="utf-8",
    )
    try:
        _run_chrome(
            [
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                "--hide-scrollbars",
                f"--window-size=512,512",
                f"--screenshot={output_path}",
                temp_html.resolve().as_uri(),
            ]
        )
    finally:
        if temp_html.exists():
            temp_html.unlink()
    return output_path


def export_pdf(html_path: Path, pdf_path: Path) -> Path:
    _run_chrome(
        [
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            "--hide-scrollbars",
            "--run-all-compositor-stages-before-draw",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ]
    )
    return pdf_path


def _set_cell_background(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _add_hyperlink(paragraph: Any, url: str, text: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "EA580C")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def export_docx(context: dict[str, Any], docx_path: Path, logo_path: Path, qr_path: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    header = document.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(1.0)
    header.columns[1].width = Inches(6.8)
    header_cell_logo = header.rows[0].cells[0]
    header_cell_text = header.rows[0].cells[1]
    header_cell_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    header_cell_text.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    header_cell_logo.paragraphs[0].add_run().add_picture(str(logo_path), width=Inches(0.72))
    heading = header_cell_text.paragraphs[0]
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading_run = heading.add_run("OxSteed Consulting")
    heading_run.font.name = "Arial"
    heading_run.font.bold = True
    heading_run.font.size = Pt(20)
    heading_run.font.color.rgb = RGBColor(0x16, 0x21, 0x3E)
    subheading = header_cell_text.add_paragraph(context["tagline"])
    subheading.paragraph_format.space_before = Pt(2)
    subheading.paragraph_format.space_after = Pt(0)
    subheading.runs[0].font.size = Pt(10)
    subheading.runs[0].font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    intro = document.add_paragraph()
    intro.paragraph_format.space_before = Pt(10)
    intro.paragraph_format.space_after = Pt(8)
    intro_run = intro.add_run(context["subtitle"])
    intro_run.font.size = Pt(10)
    intro_run.font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    def add_section_title(title: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(title)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x16, 0x21, 0x3E)

    add_section_title("What We Offer")
    offer_table = document.add_table(rows=1, cols=3)
    offer_table.autofit = True
    for index, card in enumerate(context["feature_cards"]):
        cell = offer_table.rows[0].cells[index]
        _set_cell_background(cell, "FFF8EF")
        title = cell.paragraphs[0]
        title_run = title.add_run(card["title"])
        title_run.font.bold = True
        title_run.font.size = Pt(11)
        title_run.font.color.rgb = RGBColor(0x16, 0x21, 0x3E)
        body = cell.add_paragraph(card["body"])
        body.paragraph_format.space_before = Pt(2)
        body.runs[0].font.size = Pt(9)
        body.runs[0].font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    add_section_title("How It Works")
    for step in context["how_it_works"]:
        paragraph = document.add_paragraph(style=document.styles["Normal"])
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f'{step["step"]}  {step["title"]}: ')
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)
        body_run = paragraph.add_run(step["body"])
        body_run.font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    add_section_title("Consulting Packages")
    package_table = document.add_table(rows=1, cols=3)
    package_table.autofit = True
    for index, package in enumerate(context["packages"]):
        cell = package_table.rows[0].cells[index]
        if package["popular"]:
            _set_cell_background(cell, "FFF5EB")
        name = cell.paragraphs[0]
        name_run = name.add_run(f'{package["name"]} - {package["price"]}')
        name_run.font.bold = True
        name_run.font.size = Pt(11)
        name_run.font.color.rgb = RGBColor(0x16, 0x21, 0x3E)

        timeline = cell.add_paragraph(package["timeline"])
        timeline.runs[0].font.size = Pt(9)
        timeline.runs[0].font.bold = True
        timeline.runs[0].font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

        for_who = cell.add_paragraph(package["for_who"])
        for_who.runs[0].font.size = Pt(8.5)
        for_who.runs[0].font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

        for feature in package["highlights"]:
            bullet = cell.add_paragraph(style=document.styles["Normal"])
            bullet.paragraph_format.left_indent = Inches(0.16)
            bullet.paragraph_format.space_after = Pt(0)
            bullet_run = bullet.add_run(f"• {feature}")
            bullet_run.font.size = Pt(8.5)
            bullet_run.font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    add_section_title("Why OxSteed?")
    for point in context["why_oxsteed"]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"• {point}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    document.add_section(WD_SECTION_START.CONTINUOUS)
    cta_table = document.add_table(rows=1, cols=2)
    cta_table.autofit = False
    cta_table.columns[0].width = Inches(5.8)
    cta_table.columns[1].width = Inches(1.4)
    left = cta_table.rows[0].cells[0]
    right = cta_table.rows[0].cells[1]
    _set_cell_background(left, "FFF7ED")
    _set_cell_background(right, "FFF7ED")
    left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    cta_title = left.paragraphs[0]
    cta_title_run = cta_title.add_run(context["cta_title"])
    cta_title_run.font.bold = True
    cta_title_run.font.size = Pt(15)
    cta_title_run.font.color.rgb = RGBColor(0x16, 0x21, 0x3E)

    cta_body = left.add_paragraph(context["cta_body"])
    cta_body.runs[0].font.size = Pt(9)
    cta_body.runs[0].font.color.rgb = RGBColor(0x55, 0x61, 0x7B)
    cta_link = left.add_paragraph()
    _add_hyperlink(cta_link, context["cta_url"], context["cta_label"])

    right_paragraph = right.paragraphs[0]
    right_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    right_paragraph.add_run().add_picture(str(qr_path), width=Inches(1.1))
    qr_caption = right.add_paragraph("Scan to apply")
    qr_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_caption.runs[0].font.size = Pt(8)
    qr_caption.runs[0].font.bold = True
    qr_caption.runs[0].font.color.rgb = RGBColor(0x55, 0x61, 0x7B)

    document.save(docx_path)
    return docx_path


def write_outputs(output_dir: Path = OUTPUT_DIR) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    context = build_context()
    html_path = output_dir / HTML_OUTPUT.name
    pdf_path = output_dir / PDF_OUTPUT.name
    docx_path = output_dir / DOCX_OUTPUT.name
    qr_path = output_dir / QR_OUTPUT.name
    logo_png_path = output_dir / LOGO_PNG_OUTPUT.name

    _write_file(qr_path, _make_qr_png_bytes(CTA_URL))
    _write_file(html_path, render_html(context))
    _rasterize_logo_png(logo_png_path)
    export_pdf(html_path, pdf_path)
    export_docx(context, docx_path, logo_png_path, qr_path)

    return {
        "html": html_path,
        "pdf": pdf_path,
        "docx": docx_path,
        "qr": qr_path,
        "logo_png": logo_png_path,
    }


def main() -> None:
    outputs = write_outputs()
    for label, path in outputs.items():
        print(f"{label}: {path}")


if __name__ == "__main__":
    main()
